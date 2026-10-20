from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APIClient

from events.models import Event
from users.models import User

from .models import Expense


def make_user(n, role='organizer'):
    return User.objects.create_user(
        email=f'user{n}@test.com', username=f'user{n}', password='Pass1234',
        first_name='User', last_name=str(n), role=role,
    )


class ExpenseApiTests(TestCase):
    def setUp(self):
        self.organizer = make_user(1)
        self.other = make_user(2)
        self.client = APIClient()
        self.client.force_authenticate(self.organizer)
        self.event = Event.objects.create(
            title='Conferencia', event_date='2026-12-01',
            organizer=self.organizer, budget=Decimal('10000'),
        )
        self.other_event = Event.objects.create(
            title='Ajeno', event_date='2026-12-01', organizer=self.other,
        )
        self.list_url = reverse('expense-list')

    def _create(self, **kwargs):
        data = {'event': self.event.id, 'concept': 'Catering', 'category': 'catering', 'amount': '4500'}
        data.update(kwargs)
        return self.client.post(self.list_url, data, format='json')

    # --- Alta y validaciones ---

    def test_create_expense(self):
        response = self._create(paid_by='Juan Díaz')

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['category_label'], 'Catering')
        self.assertEqual(Expense.objects.count(), 1)

    def test_amount_must_be_positive(self):
        for amount in ('0', '-50'):
            response = self._create(amount=amount)
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, amount)
        self.assertEqual(Expense.objects.count(), 0)

    def test_concept_required(self):
        response = self._create(concept='   ')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_event_required(self):
        response = self.client.post(
            self.list_url, {'concept': 'X', 'amount': '10'}, format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    # --- Permisos ---

    def test_cannot_add_expense_to_other_users_event(self):
        response = self._create(event=self.other_event.id)

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(Expense.objects.count(), 0)

    def test_cannot_see_or_edit_other_users_expense(self):
        ajeno = Expense.objects.create(event=self.other_event, concept='X', amount=10)
        url = reverse('expense-detail', args=[ajeno.id])

        self.assertEqual(self.client.get(url).status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            self.client.patch(url, {'amount': '1'}, format='json').status_code,
            status.HTTP_403_FORBIDDEN,
        )
        self.assertEqual(self.client.delete(url).status_code, status.HTTP_403_FORBIDDEN)

    def test_list_only_shows_own_expenses(self):
        Expense.objects.create(event=self.event, concept='Propio', amount=10)
        Expense.objects.create(event=self.other_event, concept='Ajeno', amount=10)

        response = self.client.get(self.list_url)

        self.assertEqual([e['concept'] for e in response.data['results']], ['Propio'])

    def test_collaborator_cannot_access(self):
        client = APIClient()
        client.force_authenticate(make_user(3, role='collaborator'))
        self.assertEqual(client.get(self.list_url).status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_sees_everything(self):
        Expense.objects.create(event=self.other_event, concept='Ajeno', amount=10)
        client = APIClient()
        client.force_authenticate(make_user(4, role='admin'))

        self.assertEqual(client.get(self.list_url).data['count'], 1)

    # --- Edicion y borrado ---

    def test_update_and_delete(self):
        expense_id = self._create().data['id']
        url = reverse('expense-detail', args=[expense_id])

        response = self.client.patch(url, {'amount': '5000'}, format='json')
        self.assertEqual(Decimal(response.data['amount']), Decimal('5000'))

        self.assertEqual(self.client.delete(url).status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Expense.objects.filter(id=expense_id).exists())

    def test_expense_cannot_be_moved_to_another_event(self):
        expense_id = self._create().data['id']
        self.client.patch(
            reverse('expense-detail', args=[expense_id]),
            {'event': self.other_event.id}, format='json',
        )
        self.assertEqual(Expense.objects.get(id=expense_id).event_id, self.event.id)

    def test_filter_by_event(self):
        second = Event.objects.create(title='Otro', event_date='2026-12-02', organizer=self.organizer)
        Expense.objects.create(event=self.event, concept='A', amount=10)
        Expense.objects.create(event=second, concept='B', amount=10)

        response = self.client.get(self.list_url, {'event': second.id})

        self.assertEqual([e['concept'] for e in response.data['results']], ['B'])


class ExpenseSummaryTests(TestCase):
    def setUp(self):
        self.organizer = make_user(1)
        self.client = APIClient()
        self.client.force_authenticate(self.organizer)
        self.url = reverse('expense-summary')

    def _event(self, budget):
        return Event.objects.create(
            title='E', event_date='2026-12-01', organizer=self.organizer, budget=budget,
        )

    def test_summary_within_budget(self):
        event = self._event(Decimal('10000'))
        Expense.objects.create(event=event, concept='Catering', category='catering', amount=4500)
        Expense.objects.create(event=event, concept='Bebidas', category='catering', amount=500)
        Expense.objects.create(event=event, concept='Salón', category='venue', amount=2800)

        data = self.client.get(self.url, {'event': event.id}).data

        self.assertEqual(data['total'], Decimal('7800'))
        self.assertEqual(data['remaining'], Decimal('2200'))
        self.assertEqual(data['percentage_used'], 78.0)
        self.assertFalse(data['over_budget'])
        self.assertEqual(data['count'], 3)
        self.assertEqual(data['by_category'][0]['category'], 'catering')
        self.assertEqual(data['by_category'][0]['total'], Decimal('5000'))

    def test_summary_over_budget(self):
        event = self._event(Decimal('1000'))
        Expense.objects.create(event=event, concept='X', amount=1500)

        data = self.client.get(self.url, {'event': event.id}).data

        self.assertTrue(data['over_budget'])
        self.assertEqual(data['remaining'], Decimal('-500'))

    def test_summary_without_budget(self):
        event = self._event(None)
        Expense.objects.create(event=event, concept='X', amount=100)

        data = self.client.get(self.url, {'event': event.id}).data

        self.assertIsNone(data['budget'])
        self.assertIsNone(data['remaining'])
        self.assertFalse(data['over_budget'])

    def test_summary_empty_event(self):
        event = self._event(Decimal('500'))
        data = self.client.get(self.url, {'event': event.id}).data
        self.assertEqual(data['total'], 0)
        self.assertEqual(data['by_category'], [])

    def test_summary_of_other_users_event_is_forbidden(self):
        other_event = Event.objects.create(
            title='Ajeno', event_date='2026-12-01', organizer=make_user(2),
        )
        response = self.client.get(self.url, {'event': other_event.id})
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_dashboard_includes_total_expenses(self):
        event = self._event(None)
        Expense.objects.create(event=event, concept='X', amount=100)
        Expense.objects.create(event=event, concept='Y', amount=250)

        data = self.client.get(reverse('dashboard-stats')).data

        self.assertEqual(data['total_expenses'], Decimal('350'))


class EventBudgetTests(TestCase):
    def setUp(self):
        self.organizer = make_user(1)
        self.client = APIClient()
        self.client.force_authenticate(self.organizer)

    def test_create_event_with_budget(self):
        response = self.client.post(reverse('event-list'), {
            'title': 'Con presupuesto', 'event_date': '2026-12-01', 'budget': '15000',
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(Event.objects.get().budget, Decimal('15000'))

    def test_negative_budget_rejected(self):
        response = self.client.post(reverse('event-list'), {
            'title': 'X', 'event_date': '2026-12-01', 'budget': '-1',
        }, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class ExpenseExportTests(TestCase):
    def setUp(self):
        self.organizer = make_user(1)
        self.other = make_user(2)
        self.client = APIClient()
        self.client.force_authenticate(self.organizer)
        self.event = Event.objects.create(
            title='Asado de fin de año', event_date='2026-12-01',
            organizer=self.organizer, budget=Decimal('10000'),
        )
        Expense.objects.create(event=self.event, concept='Carne', category='catering',
                               amount=Decimal('4500.5'), paid_by='Eloy', expense_date='2026-11-20')
        Expense.objects.create(event=self.event, concept='Parlante', category='sound_lighting',
                               amount=Decimal('1200'), expense_date='2026-11-10')

    def _sheet(self):
        import io

        from openpyxl import load_workbook
        response = self.client.get(reverse('expense-export'), {'event': self.event.id})
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(io.BytesIO(response.content)).active

    def test_export_expenses_to_excel(self):
        import datetime
        response, ws = self._sheet()

        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertIn('gastos-asado-de-fin-de-ano-', response['Content-Disposition'])
        self.assertTrue(response['Content-Disposition'].rstrip('"').endswith('.xlsx'))
        self.assertEqual([c.value for c in ws[1]], ['Fecha', 'Concepto', 'Categoría', 'Pagado por', 'Monto', 'Notas'])
        # Ordenados por fecha; la fecha es una fecha real y el monto un numero con formato moneda
        self.assertEqual(ws['A2'].value, datetime.datetime(2026, 11, 10))
        self.assertEqual(ws['A2'].number_format, 'DD/MM/YYYY')
        self.assertEqual(ws['C2'].value, 'Sonido e iluminación')
        self.assertEqual(ws['E3'].value, 4500.5)
        self.assertEqual(ws['E3'].number_format, '"$" #,##0.00')
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual(ws.auto_filter.ref, 'A1:F3')

    def test_totals_are_formulas(self):
        _, ws = self._sheet()

        # Fila 4 vacia; luego total, presupuesto y disponible
        self.assertIsNone(ws['B4'].value)
        self.assertEqual(ws['B5'].value, 'Total gastado')
        self.assertEqual(ws['E5'].value, '=SUM(E2:E3)')
        self.assertTrue(ws['B5'].font.bold)
        self.assertEqual(ws['E6'].value, 10000)
        self.assertEqual(ws['E7'].value, '=E6-E5')
        self.assertEqual(ws['E7'].number_format, '"$" #,##0.00')

    def test_columns_are_wide_enough(self):
        _, ws = self._sheet()

        self.assertGreaterEqual(ws.column_dimensions['A'].width, 10)
        self.assertGreaterEqual(ws.column_dimensions['C'].width, len('Sonido e iluminación'))

    def test_export_without_budget_has_no_budget_rows(self):
        self.event.budget = None
        self.event.save()

        _, ws = self._sheet()

        self.assertNotIn('Presupuesto', [c.value for c in ws['B']])

    def test_export_without_expenses(self):
        Expense.objects.all().delete()

        _, ws = self._sheet()

        self.assertEqual(ws['B3'].value, 'Total gastado')
        self.assertEqual(ws['E3'].value, 0)

    def test_formula_like_text_is_saved_as_text(self):
        Expense.objects.create(event=self.event, concept='=HYPERLINK("http://x")', amount=Decimal('1'),
                               expense_date='2026-11-30')

        _, ws = self._sheet()

        cell = ws['B4']
        self.assertEqual(cell.value, '=HYPERLINK("http://x")')
        self.assertEqual(cell.data_type, 's')

    def test_other_organizer_cannot_export(self):
        client = APIClient()
        client.force_authenticate(self.other)

        response = client.get(reverse('expense-export'), {'event': self.event.id})

        self.assertEqual(response.status_code, 403)
