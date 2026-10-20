"""Reportes en Excel (.xlsx) con formato: moneda, fechas, anchos de columna,
encabezado fijo y filtros."""
import datetime
import io
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from django.utils.functional import Promise
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

# Excel muestra los separadores segun la configuracion regional:
# en espanol, '$ 28.500,50'
MONEY = '"$" #,##0.00'
DATE = 'DD/MM/YYYY'
DATETIME = 'DD/MM/YYYY HH:MM'

HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='4F46E5')
BOLD = Font(bold=True)


def filename_for(prefix, title):
    """('gastos', 'Asado de fin de año') -> 'gastos-asado-de-fin-de-ano-2026-10-20.xlsx'"""
    return f'{prefix}-{slugify(title) or "evento"}-{timezone.localdate():%Y-%m-%d}.xlsx'


def local_naive(value):
    """Excel no guarda zonas horarias: se escribe la hora de Argentina."""
    return timezone.localtime(value).replace(tzinfo=None) if value else None


def yes_no(value):
    return 'Sí' if value else 'No'


class Report:
    def __init__(self, sheet_title, header, formats=None):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = sheet_title
        self.formats = formats or {}  # {indice de columna: formato de numero}
        self.widths = [len(h) for h in header]
        self.ws.append(header)
        # Contador propio: una fila vacia no crea celdas y openpyxl no la cuenta
        self._row = 1
        for cell in self.ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical='center')
        self.ws.freeze_panes = 'A2'  # el encabezado queda fijo al bajar
        # Al imprimir: horizontal y todas las columnas en el ancho de la hoja
        self.ws.page_setup.orientation = 'landscape'
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.page_setup.fitToWidth = 1
        self.ws.page_setup.fitToHeight = 0

    @property
    def last_row(self):
        return self._row

    def add_row(self, values, bold=False):
        self._row += 1
        row = self._row
        for col, value in enumerate(values, start=1):
            if isinstance(value, Decimal):
                value = float(value)  # Excel guarda numeros en coma flotante
            elif isinstance(value, Promise):
                value = str(value)  # textos traducibles de Django (etiquetas)
            cell = self.ws.cell(row=row, column=col, value=value)
            if isinstance(value, str) and value.startswith('=') and not getattr(value, 'is_formula', False):
                # Texto escrito por un usuario que empieza con '=': se guarda
                # como texto para que Excel no lo ejecute como formula
                cell.data_type = 's'
            fmt = self.formats.get(col - 1)
            if fmt and value is not None:
                cell.number_format = fmt
            if bold:
                cell.font = BOLD
            self._measure(col - 1, value)
        return row

    def _measure(self, index, value):
        if value is None or getattr(value, 'is_formula', False):
            return
        if isinstance(value, datetime.datetime):
            length = 16
        elif isinstance(value, datetime.date):
            length = 10
        elif isinstance(value, float):
            length = len(f'$ {value:,.2f}')
        else:
            length = max((len(line) for line in str(value).splitlines()), default=0)
        while len(self.widths) <= index:
            self.widths.append(0)
        self.widths[index] = max(self.widths[index], length)

    def add_filter(self, last_data_row):
        cols = len(self.widths)
        self.ws.auto_filter.ref = f'A1:{get_column_letter(cols)}{max(last_data_row, 1)}'

    def response(self, filename):
        for i, width in enumerate(self.widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = min(width + 3, 60)
        buf = io.BytesIO()
        self.wb.save(buf)
        response = HttpResponse(buf.getvalue(), content_type=XLSX_TYPE)
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class Formula(str):
    """Formula de Excel escrita por el sistema (no por un usuario)."""
    is_formula = True
